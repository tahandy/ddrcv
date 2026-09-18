"""Build three complete, content-identical 10-slide operator guide designs.

Uses the same document libraries as build_operator_deck.py. Native editable
PowerPoint objects and a companion PDF share a measured layout specification.
"""
from pathlib import Path
import argparse
import hashlib
import json
import math
import re
import statistics
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from pypdf import PdfReader
import pypdfium2 as pdfium
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'doc/slides/variants'
ASSETS = OUT / 'assets'
W, H = 13.333333, 7.5

THEMES = [
    dict(id='a_field_manual', letter='A', name='Field manual',
         description='Light, practical, clear hierarchy. Best default for training and reference.',
         bg='F5F7FA', panel='FFFFFF', ink='172B46', muted='53647A', rule='D6DEE8',
         accent='205CB0', good='246E5B', warning='955215', tint='E5EDF8',
         title='UIB', label='UIB', x=.6, width=12.13, top=1.78),
    dict(id='b_control_room', letter='B', name='Control room',
         description='Dark dashboard, cyan cues, persistent step rail. Best for a screen beside the station.',
         bg='111920', panel='1C2933', ink='EFF5F7', muted='B2C2CC', rule='3B4D5A',
         accent='70DEEB', good='B0DC89', warning='F2C686', tint='243C46',
         title='UIB', label='MonoB', x=1.16, width=11.55, top=1.78),
    dict(id='c_editorial', letter='C', name='Editorial',
         description='Warm paper, serif headlines, open columns and fine rules. Best for a guided workshop.',
         bg='F7F1E7', panel='F0E6D7', ink='302D29', muted='685E53', rule='CCBDA9',
         accent='A6402F', good='47634C', warning='865416', tint='EBDED0',
         title='Serif', label='UIB', x=.7, width=11.93, top=1.9),
]

NOTES = [
    '''Use the organizer's approved ZIP or checkout. Extract completely to a writable
folder such as C:/Games/ddrcv, outside Program Files. Setup option 1 installs local
x64 CPython 3.12; no existing Python is required. Options 2/3 select a compatible
existing interpreter. Option 4 reuses an environment on subsequent runs. Setup checks
actual OCR and WebView2. Retain error text on failure. Keep .venv and .tools/python
with the installation; run setup again on another computer. The renderer repository
and its assets are separate. Source: doc/installation.md; operator_walkthrough.md §§1–2.
Screenshots in this deck are real forms captured in an isolated stopped session with
example paths, no credentials, and no hardware or external integrations started.''',
    '''The UI detail is intentionally limited to source and mode controls. Capture
camera: Scan, select the device, verify index. Different computers can assign different
indices. RTSP requires full 1280x720 frames. Local video may explicitly resize full
16:9 footage. Cropped or overlaid broadcasts remain unsuitable. Still-image diagnostics
cannot establish temporal confirmation. EX is the desktop default; browser workbench
defaults to Money. An existing profile can override either. A wrong mode affects live
filtering and results interpretation. Leave Custom filter to a maintainer. Keep default
publisher localhost:9000 for one machine; only one recognition session may own that
port/outbox. Source: operator_walkthrough.md §3; driver_config.py.''',
    '''The screenshots show the stopped UI as a control-location reference, not an
observed live session. Start uses the current form. Watch Running and healthy capture.
View stream fetches a single snapshot; Refresh frame fetches another. Fresh capture
does not imply fresh accepted scores; inspect the accepted-read age independently.
Beginning before song splash provides presence evidence. Brief state confirmation delays
are normal. Compare final scores, including full-combo endings, against the machine.
Use the separate browser workbench for pause, seek, slow playback, and diagnostic export.
Source: operator_walkthrough.md §4; tournament_desktop.md; tournament_handoff.md.''',
    '''Both integrations are optional. Renderer: obtain the separate ddr_score_render
repository/assets, select its root (containing apps/tbd5.py), apply path, then Start or
Launch renderer. The renderer's visible output must match accepted scores; a process
or client count is insufficient. Capture the window in OBS. OBS: enable its WebSocket
server, enter host/port/password (usually localhost:4455), apply, connect and Check scenes.
Default map: song_select/login/total_result to Commentator; song_splash/song_playing/
song_result to Gameplay. Use exact existing scene names; null/omitted states hold and
unknown always holds. No scenes are created. Manual switching inside DDRCV pauses
automation; Enable automation resumes. Pause automation before switching directly in
OBS. Set automatic launch/connect only after rehearsal. Source: walkthrough §§5–6;
tournament_integrations.py; publish/obs_policy.py.''',
    '''Capture results defaults on for new profiles; a saved off setting is preserved.
Choose a writable Screenshot folder. Default exported copies are in screenshots, while
durable originals/history are in .workbench/results unless overridden. Filenames use UTC
and a unique suffix. Include parsed names in filenames affects naming only. Screenshots
only and Read names on results apply immediately; Apply results settings applies other
Results controls without restarting gameplay or integrations. Saved results are unchanged.
Parsed scores can post first and later names update the same message. Screenshots only
saves a detected episode on departure or its original 10-second deadline and skips results
score/name OCR; gameplay scoring continues. Money and screenshot-only results do not
declare an EX winner. The separate login/song-splash names switch controls live names.
Source: walkthrough §7; results_publishing.md; tournament_desktop.md.''',
    '''Test Discord Webhook sends a real message using the entered URL, even when
capture and publishing are off. It does not apply settings or enable publishing. Verify
the intended channel. Publish results to Discord also checks Capture results; apply the
Results controls. Retry/send needs Running capture, Capture results, Discord publishing,
and a webhook. Inspect screenshots before sending. Queued/Sending/Retrying are ongoing;
after failure fix the cause and Retry delivery. Retry name update edits the original
message. Republish intentionally creates another post. Disabling publishing and applying
pauses queues; an in-flight send may finish. Re-enabling resumes pending work. An entirely
missed results screen cannot be recovered automatically. Source: operator_reference.md,
Results recovery; desktop/static/desk.js; results_publishing.md.''',
    '''Open loads a form without starting or replacing the active configuration. Save
writes the form to a profile, including actual secrets/absolute paths, and does not apply
pending edits. Start uses the current form. Apply configuration while stopped prepares
it; Start runs it. Apply and Restart while running asks before resetting tracking.
Results-only Apply preserves live gameplay and integrations. Name/screenshot-only toggles
are immediate. Stop stops recognition and owned integrations but leaves the application
available. Window close keeps the backend running; reopen from the tray or launcher.
Exit application stops the backend and owned renderer; it does not close OBS itself.
Source: operator_reference.md, Controls; desktop/controller.py and static/desk.js.''',
    '''Rehearse on the actual event computer, capture device, renderer and scene
collection. Observe selection, splash, gameplay, song results and total results. Check
single-player presence, full-combo final score, capture reconnect and manual OBS/resume.
Verify actual screenshot files and Discord delivery. Installation checks and these
illustrative UI screenshots do not establish hardware readiness. Daily startup uses
the saved profile, checks device/mode/integrations, then explicit Start. A reopened panel
can already have a running backend; do not start a second app/CLI/workbench on its port.
Source: walkthrough §8; operator_reference.md, Start an event session.''',
    '''Diagnose capture freshness separately from accepted-score freshness. Stale capture:
check connection, index/RTSP, native 720p, ownership and errors. Frozen requested preview
alone requires Refresh frame. Starting mid-song can lack presence evidence until a splash.
For a bad high accepted score first verify mode/input, then reset only that player's
tracking under Read details & recent errors. This cannot correct saved results. A camera
handoff probes the new source before replacing the old; success resets tracking while
preserving integrations/results. A failed probe retains the old feed. For hardware that
cannot open both, Stop/Start between songs. Port/outbox conflicts require stopping the
other session; closing its panel is insufficient. Screenshot export errors may leave
internal evidence still available for Discord. Source: operator_reference.md.''',
    '''Record source folder/version, profile path, device/index/resolution, score mode,
renderer path, OBS scene collection/mapping, screenshot folder, Discord channel owner,
technical contact and manual evidence procedure. Fill the walkthrough's event table.
Do not put passwords or webhook tokens in the deck. Profile JSON contains real secrets
despite masking in the UI and saves absolute paths, which need checking on another PC.
At shutdown inspect unresolved delivery, save intended changes, then Exit application.
Keep screenshot exports and the durable outbox for evidence. Desktop logs/settings live
in %LOCALAPPDATA%/DDRCV/tournament. Before updates fully exit, update source, rerun setup,
reuse a compatible environment, and rehearse. Detailed references: operator_walkthrough.md,
operator_reference.md, installation.md, tournament_handoff.md, results_publishing.md.''',
]


class Deck:
    def __init__(self, theme):
        self.t = theme
        self.slides = []
        self.x, self.w, self.top = theme['x'], theme['width'], theme['top']

    def box(self, s, x, y, w, h, fill, stroke=None, radius=False):
        s['elements'].append(dict(kind='box',x=x,y=y,w=w,h=h,fill=fill,stroke=stroke,radius=radius))

    def txt(self, s, value, x, y, w, size=20, color=None, font='UI', max_h=None, role='content'):
        lines=[]
        for para in value.split('\n'):
            current=''
            for word in para.split():
                test=(current+' '+word).strip()
                if current and pdfmetrics.stringWidth(test,font,size)>w*72:
                    lines.append(current);current=word
                else: current=test
            lines.append(current)
        leading=size*1.18
        h=len(lines)*leading/72+.045
        assert max_h is None or h<=max_h+.003, (self.t['id'],s['index'],value,h,max_h)
        assert all(pdfmetrics.stringWidth(l,font,size)<=w*72+1 for l in lines),value
        s['elements'].append(dict(kind='text',x=x,y=y,w=w,h=h,lines=lines,value=value,
            size=size,leading=leading,font=font,color=color or self.t['ink'],role=role))
        if role=='content':s['copy'].append(value)
        return h

    def rule(self,s,x,y,w,color=None):
        self.box(s,x,y,w,.012,color or self.t['rule'])

    def new(self,section,title,sub):
        t=self.t;n=len(self.slides)+1
        s=dict(index=n,title=title,section=section,notes=NOTES[n-1],elements=[],copy=[])
        self.slides.append(s)
        if t['letter']=='A':
            self.box(s,0,0,W,.1,t['accent'])
            self.txt(s,'DDRCV / OPERATOR GUIDE',self.x,.3,6,11,t['accent'],'UIB',role='chrome')
            self.txt(s,f'{section.upper()}   /   {n:02d}',10.1,.3,2.6,11,t['muted'],'UIB',role='chrome')
            ty=.67
        elif t['letter']=='B':
            self.box(s,0,0,.78,H,t['panel'])
            for i in range(10):
                yy=1.26+i*.5
                if i+1==n:self.box(s,.13,yy-.03,.52,.35,t['accent'])
                self.txt(s,f'{i+1:02d}',.24,yy,.36,11,t['bg'] if i+1==n else t['muted'],'MonoB',role='chrome')
            self.txt(s,'DDRCV  /  '+section.upper(),self.x,.3,10.4,12,t['accent'],'MonoB',role='chrome')
            self.box(s,12.05,.26,.65,.3,t['tint'])
            self.txt(s,'GUIDE',12.12,.3,.52,9,t['accent'],'MonoB',role='chrome')
            ty=.67
        else:
            self.txt(s,'DDRCV   /   '+section.upper(),self.x,.3,10.5,11,t['accent'],'UIB',role='chrome')
            self.txt(s,f'{n:02d}',11.59,.25,1.06,36,t['accent'],'Serif',role='chrome')
            self.rule(s,self.x,.74,self.w)
            ty=.89
        self.txt(s,title,self.x,ty,self.w-(.1 if t['letter']!='C' else .2),30,
                 font=t['title'],max_h=.59)
        self.txt(s,sub,self.x,ty+.61,self.w,16,t['muted'],max_h=.56)
        self.rule(s,self.x,7.07,self.w)
        self.txt(s,f'DDRCV  •  {t["name"]}  •  18 Sep 2026',self.x,7.18,self.w-1,9,t['muted'],role='chrome')
        self.txt(s,f'{n:02d} / 10',self.x+self.w-.8,7.16,.8,10,t['muted'],'UIB',role='chrome')
        return s

    def panel(self,s,x,y,w,h):
        t=self.t
        if t['letter']=='C':self.rule(s,x,y,w,t['accent'])
        else:self.box(s,x,y,w,h,t['panel'],t['rule'] if t['letter']=='A' else None,t['letter']=='A')

    def heading(self,s,value,x,y,w,size=20):
        return self.txt(s,value,x,y,w,size,self.t['accent'],'UIB')

    def note(self,s,label,body,y=6.34,warning=False):
        t=self.t
        self.box(s,self.x,y,self.w,.54,t['tint'])
        self.box(s,self.x,y,.042,.54,t['warning'] if warning else t['good'])
        lh=self.txt(s,label,self.x+.17,y+.13,1.54,12,t['warning'] if warning else t['good'],t['label'],max_h=.38)
        self.txt(s,body,self.x+1.93,y+.12,self.w-2.1,16,max_h=.4)

    def image(self,s,name,x,y,w,max_h=None):
        path=ASSETS/name
        with Image.open(path) as im:ratio=im.width/im.height
        h=w/ratio
        if max_h and h>max_h:h,w=max_h,max_h*ratio
        s['elements'].append(dict(kind='image',path=str(path),x=x,y=y,w=w,h=h))
        return h

    def caption(self,s,x,y,w):
        self.txt(s,'UI reference: stopped session; example settings.',x,y,w,10,self.t['muted'],role='chrome')

    def step(self,s,n,title,body,x,y,w,size=19):
        t=self.t
        if t['letter']=='C':
            self.txt(s,str(n),x,y-.04,.39,26,t['accent'],'Serif')
        else:
            self.box(s,x,y,.34,.34,t['tint'])
            self.txt(s,str(n),x+.09,y+.025,.2,14,t['accent'],'UIB')
        self.txt(s,title,x+.51,y,w-.55,size,t['ink'],'UIB')
        self.txt(s,body,x+.51,y+.4,w-.57,size-2,max_h=.85)

    def table_row(self,s,y,left,right,height=.68,split=3.4,size=19):
        t=self.t
        if t['letter']=='B':self.box(s,self.x,y,self.w,height-.06,t['panel'])
        else:self.rule(s,self.x,y+height-.08,self.w)
        self.txt(s,left,self.x+.12,y+.1,split-.23,size,t['accent'],'UIB',max_h=height-.12)
        self.txt(s,right,self.x+split+.1,y+.1,self.w-split-.25,size,max_h=height-.12)

    def content(self):
        x,w,y=self.x,self.w,self.top;t=self.t
        s=self.new('Setup','Install DDRCV and open the panel','A Windows guide from first setup to event handoff.')
        # A single compact flow replaces the old cover and separate architecture slide.
        cells=[('GAME FEED','1280 × 720'),('DDRCV','Read state + scores'),('OUTPUTS','Renderer / OBS • screenshots / Discord')]
        widths=[w*.25,w*.25,w*.45]
        xx=x
        for i,((a,b),ww) in enumerate(zip(cells,widths)):
            self.box(s,xx,y,ww,.85,t['tint'])
            self.txt(s,a,xx+.15,y+.12,ww-.3,12,t['accent'],t['label'])
            self.txt(s,b,xx+.15,y+.4,ww-.3,17)
            if i<2:self.txt(s,'→',xx+ww+.035,y+.22,.3,21,t['accent'])
            xx+=ww+w*.025
        left=w*.295;right_x=x+left+.4;right_w=w-left-.4
        self.heading(s,'Before setup',x,y+1.17,left)
        self.txt(s,'64-bit Windows\nFirst-install internet\nFull, clean game feed',x,y+1.68,left,21)
        self.txt(s,'Optional: event profile,\nrenderer, OBS, webhook.',x,y+3.04,left,18,t['muted'])
        for i,(a,b) in enumerate([
            ('Extract',r'Put DDRCV in a writable folder, e.g. C:\Games\ddrcv.'),
            ('Set up','setup_desktop.bat → option 1 (local Python).'),
            ('Wait','Continue after “Setup complete”.'),
            ('Launch','launch_desktop.vbs → the panel opens stopped.')]):
            yy=y+1.13+i*.79
            self.txt(s,f'{i+1:02d}',right_x,yy,.5,19,t['accent'],t['label'])
            self.txt(s,a,right_x+.61,yy,1.15,19,font='UIB')
            self.txt(s,b,right_x+1.94,yy,right_w-1.94,18,max_h=.68)
            if i<3:self.rule(s,right_x,yy+.69,right_w)
        self.note(s,'SETUP CHECK','No Python install needed beforehand. Wait for success before launching.')

        s=self.new('Setup','Select the feed and the score mode','Status & configuration → Source and Score mode')
        ph=self.image(s,'source-controls.png',x,y,w,1.58)
        self.caption(s,x,y+ph+.08,w)
        cy=y+ph+.48;cw=(w-.54)/3
        for i,(title,body) in enumerate([
            ('Capture camera','Scan → select device.\nVerify Camera index.'),
            ('RTSP stream','Enter Stream URL.\nFeed must be 1280 × 720.'),
            ('Local video','Browse to a recording.\nResize only full 16:9 files.')]):
            xx=x+i*(cw+.27)
            self.heading(s,title,xx,cy,cw)
            self.txt(s,body,xx,cy+.43,cw,19)
        my=5.04
        self.panel(s,x,my,w,.96)
        self.txt(s,'EX',x+.16,my+.16,.5,22,t['accent'],'UIB')
        self.txt(s,'Event EX scoring.\nDesktop default.',x+.83,my+.13,w*.4-1,18)
        self.txt(s,'MONEY',x+w*.5,my+.16,1.05,19,t['accent'],'UIB')
        self.txt(s,'Score toward 1,000,000.\nWorkbench default.',x+w*.5+1.21,my+.13,w*.5-1.37,18)
        self.note(s,'CHECK THE FEED','No borders, cropping, overlays, or stretching. Check mode on every profile.',warning=True)

        s=self.new('First run','Start, then verify picture, state, and scores','Click Start → Running. Begin before the song splash for player presence.')
        ph=self.image(s,'score-readout.png',x,y,w,1.85)
        hy=y+ph+.13
        ph2=self.image(s,'frame-health.png',x,hy,w,.88)
        self.caption(s,x,hy+ph2+.07,w)
        cy=hy+ph2+.36;cw=(w-.54)/3
        for i,(a,b) in enumerate([
            ('Picture','Refresh frame for a new picture.\nCheck capture age.'),
            ('State','Splash → gameplay → results.\nConfirmed state should follow.'),
            ('Scores','Compare with the machine.\nCheck last-read age.')]):
            xx=x+i*(cw+.27);self.heading(s,f'{i+1}  {a}',xx,cy,cw)
            self.txt(s,b,xx,cy+.43,cw,18,max_h=1.0)
        self.note(s,'PASS WHEN','The picture is fresh and final scores agree, including a full-combo ending.')

        s=self.new('Optional integrations','Connect the display and OBS scene control','Apply edited paths and connection settings before using their buttons.')
        cw=(w-.54)/2
        for xx,title,sub in [(x,'Score display','Renderer tab'),(x+cw+.54,'Scene control','OBS tab')]:
            self.panel(s,xx,y,cw,4.3)
            self.heading(s,title,xx+.18,y+.18,cw-.36,23)
            self.txt(s,sub,xx+.18,y+.62,cw-.36,12,t['muted'],t['label'])
        self.step(s,1,'Select the repository','Browse to the separate ddr_score_render folder.',x+.18,y+1.02,cw-.36,19)
        self.step(s,2,'Apply, then launch','Start, or use Launch renderer after capture starts.',x+.18,y+2.1,cw-.36,19)
        self.step(s,3,'Check the actual window','Compare scores, then capture that window in OBS.',x+.18,y+3.18,cw-.36,19)
        rx=x+cw+.72
        self.step(s,1,'Connect and validate','Enable OBS WebSocket. Enter host/port/password; Apply → Connect.',rx,y+1.02,cw-.36,19)
        self.step(s,2,'Map existing scenes','Use exact scene names → Check scenes → Enable automation.',rx,y+2.1,cw-.36,19)
        self.step(s,3,'Practice manual control','Switch scene & pause automation. Enable automation resumes.',rx,y+3.18,cw-.36,19)
        self.note(s,'LIVE CHECK','Watch the renderer and OBS themselves; “connected” alone is not proof.')

        s=self.new('Results','Choose the results evidence you need','Results → Capture results ON → choose a writable Screenshot folder')
        cols=[0,w*.245,w*.45,w*.655]
        widths=[w*.245,w*.205,w*.205,w*.345]
        labels=['KEEP','SCREENSHOTS ONLY','READ NAMES','RESULT']
        self.box(s,x,y,w,.47,t['tint'])
        for xx,ww,label in zip(cols,widths,labels):
            self.txt(s,label,x+xx+.12,y+.13,ww-.2,12,t['accent'],t['label'],max_h=.32)
        modes=[('Scores + names','OFF','ON','Scores first; names can update later.'),
               ('Scores + P1/P2','OFF','OFF','Keep scores without player-name parsing.'),
               ('Screenshot only','ON','Either','Save on results departure or after 10 seconds.')]
        for i,vals in enumerate(modes):
            yy=y+.57+i*.86
            if t['letter']=='B':self.box(s,x,yy,w,.78,t['panel'])
            else:self.rule(s,x,yy+.77,w)
            for j,(xx,ww,val) in enumerate(zip(cols,widths,vals)):
                self.txt(s,val,x+xx+.12,yy+.12,ww-.22,18,t['accent'] if j==0 else t['ink'],
                         'UIB' if j<3 else 'UI',max_h=.67)
        cy=y+3.28
        self.heading(s,'Apply results settings',x,cy,w*.49,21)
        self.txt(s,'Applies Results controls without resetting gameplay scores.',x,cy+.44,w*.47,19,max_h=.76)
        self.heading(s,'Immediate switches',x+w*.53,cy,w*.47,21)
        self.txt(s,'Screenshots only and Read names on results apply as you toggle them.',x+w*.53,cy+.44,w*.47,19,max_h=.76)
        self.note(s,'LIMIT','A completely missed results screen has no automatic screenshot to recover.',warning=True)

        s=self.new('Discord + recovery','Verify the destination, then publish','Results → Webhook URL. A webhook test sends a real message.')
        lw=w*.55;rx=x+lw+.35;rw=w-lw-.35
        ph=self.image(s,'discord-webhook.png',x,y,w,1.1)
        self.caption(s,x,y+ph+.06,w)
        cy=y+ph+.4
        for i,(a,b) in enumerate([
            ('Test the webhook','Click Test Discord Webhook; inspect the channel.'),
            ('Enable automatic posts','Publish results to Discord → Apply results settings.'),
            ('Verify one result','Complete a song; inspect the screenshot and Discord post.')]):
            self.step(s,i+1,a,b,x,cy+i*.98,lw,19)
        self.heading(s,'Recent results & recovery',rx,cy,rw,21)
        for i,(a,b) in enumerate([
            ('Saved','Send saved result'),
            ('Pending','Wait; check connectivity.'),
            ('Failed','Fix → Retry delivery.'),
            ('Sent','Open Discord message.'),
            ('Republish','Creates another post.')]):
            yy=cy+.61+i*.48
            self.txt(s,a,rx,yy,1.17,17,t['accent'],'UIB')
            self.txt(s,b,rx+1.31,yy,rw-1.31,17)
            if i<4:self.rule(s,rx,yy+.4,rw)
        self.note(s,'SEND / RETRY','Needs running capture, Capture results, Discord publishing, and a webhook.')

        s=self.new('Controls','Know what applies, resets, or keeps running','The form, running session, and saved profile can contain different settings.')
        data=[('Save / Save as…','Write the form to a profile. Running settings stay unchanged.'),
              ('Start','Run with the settings currently in the form.'),
              ('Apply and Restart','Apply pending edits; reset live score tracking. Use between songs.'),
              ('Apply results settings','Apply only Results controls; keep gameplay scores.'),
              ('Close window / Alt+F4','Capture continues. Reopen from the tray or launcher.'),
              ('Stop / Exit application…','Stop capture. Exit also closes the backend; OBS stays open.')]
        for i,(a,b) in enumerate(data):self.table_row(s,y+i*.69,a,b,.65,3.76,18)
        self.note(s,'WHILE STOPPED','Apply configuration prepares settings; click Start to run.')

        s=self.new('Event readiness','Rehearse once; repeat the startup checks','Use the actual tournament computer, capture device, renderer, and OBS collection.')
        self.box(s,x,y,w,.55,t['tint'])
        self.txt(s,'Song selection → splash → gameplay → song results → total results',x+.18,y+.13,w-.36,20,t['accent'],'UIB')
        cw=(w-.56)/2
        for xx,title in [(x,'Before the event'),(x+cw+.56,'Every session')]:self.heading(s,title,xx,y+.86,cw,23)
        left=[('Scores','Two players, one player, full-combo ending.'),
              ('Recovery','Capture reconnect; manual OBS scene and resume.'),
              ('Evidence','Verify screenshot files and the actual Discord post.')]
        right=[('Open','Connect hardware → launch → open the event profile.'),
               ('Check','Source/index, EX/Money, and integrations.'),
               ('Start','Check capture freshness, accepted scores, display, and scenes.')]
        for items,xx in [(left,x),(right,x+cw+.56)]:
            for i,(a,b) in enumerate(items):self.step(s,i+1,a,b,xx,y+1.43+i*.99,cw,19)
        self.note(s,'HANDOFF TEST','Another operator can start the station and recover a saved result unaided.')

        s=self.new('Troubleshooting','Check the failing part first','Read details & recent errors explains rejected reads and recent failures.')
        data=[('No picture / stale capture','Check cable, index/URL, 720p, and capture-device ownership.'),
              ('Preview looks frozen','Refresh frame. This preview is a requested snapshot.'),
              ('Wrong / stuck score','Check framing + EX/Money; inspect rejected-read details.'),
              ('One bad accepted score','Reset P1/P2 tracking clears that live score. Saved results stay.'),
              ('Blank display / wrong scene','Check renderer output, OBS scene names, and automation.'),
              ('Port / outbox already in use','Stop the other app/workbench/CLI; it may still be running.')]
        for i,(a,b) in enumerate(data):self.table_row(s,y+i*.7,a,b,.68,3.84,18)
        self.note(s,'FALLBACK','Unreliable results text? Enable Screenshots only; gameplay scores continue.')

        s=self.new('Handoff + shutdown','Leave a usable profile and an evidence trail','Fill the event setup table in the written walkthrough before handing over.')
        cw=(w-.56)/2
        self.heading(s,'Record for the next operator',x,y,cw,22)
        for i,line in enumerate(['Profile path + DDRCV folder/version',
            'Capture device/index + EX/Money',
            'Renderer folder + OBS scenes',
            'Screenshot folder + Discord channel',
            'Technical contact + manual-results procedure']):
            yy=y+.63+i*.49
            self.box(s,x,yy+.07,.115,.115,t['accent'])
            self.txt(s,line,x+.28,yy,cw-.3,18)
        rx=x+cw+.56
        self.heading(s,'End the session',rx,y,cw,22)
        for i,(a,b) in enumerate([
            ('Inspect deliveries','Resolve failures or retain the queued evidence.'),
            ('Save intended changes','Save the profile; keep credentials private.'),
            ('Exit application…','Keep screenshots and the Results outbox.')]):
            self.step(s,i+1,a,b,rx,y+.61+i*.88,cw,19)
        by=5.32
        self.rule(s,x,by,w)
        self.txt(s,'EVIDENCE',x,by+.17,1.6,12,t['accent'],t['label'])
        self.txt(s,'screenshots/  +  .workbench/results/  (defaults)',x+1.9,by+.14,w-2,17)
        self.txt(s,'DETAILS',x,by+.61,1.6,12,t['accent'],t['label'])
        self.txt(s,'doc/operator_walkthrough.md  •  doc/operator_reference.md',x+1.9,by+.58,w-2,16)
        self.note(s,'REMEMBER','Saved profiles contain real credentials. Verify paths on another computer.',warning=True)
        return self.slides


FONT_MAP={'UI':('Segoe UI',False),'UIB':('Segoe UI',True),
          'MonoB':('Consolas',True),'Serif':('Georgia',False)}


def render(deck):
    t=deck.t;folder=OUT/t['id'];folder.mkdir(parents=True,exist_ok=True)
    slides=deck.content()
    ppt=Presentation();ppt.slide_width=Inches(W);ppt.slide_height=Inches(H)
    ppt.core_properties.title=f'DDRCV operator guide — {t["name"]}'
    ppt.core_properties.author='DDRCV'
    pdf=canvas.Canvas(str(folder/'guide.pdf'),pagesize=(W*72,H*72))
    pdf.setTitle(ppt.core_properties.title);pdf.setAuthor('DDRCV')
    overlapping=[];text_count=0
    for s in slides:
        ps=ppt.slides.add_slide(ppt.slide_layouts[6])
        ps.background.fill.solid();ps.background.fill.fore_color.rgb=RGBColor.from_string(t['bg'])
        ps.notes_slide.notes_text_frame.text=s['notes']
        pdf.bookmarkPage(str(s['index']));pdf.addOutlineEntry(f'{s["index"]:02d}. {s["title"]}',str(s['index']))
        pdf.setFillColor('#'+t['bg']);pdf.rect(0,0,W*72,H*72,fill=1,stroke=0)
        texts=[]
        for e in s['elements']:
            x,y,w,h=[e[k] for k in ('x','y','w','h')]
            assert x>=0 and y>=0 and x+w<W+.01 and y+h<H+.01,(t['id'],s['index'],e)
            if e['kind']=='text':
                text_count+=1;texts.append(e)
                shape=ps.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
                tf=shape.text_frame
                tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
                tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE
                family,bold=FONT_MAP[e['font']]
                for i,line in enumerate(e['lines']):
                    p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
                    p.font.name=family;p.font.bold=bold;p.font.size=Pt(e['size'])
                    p.font.color.rgb=RGBColor.from_string(e['color'])
                    p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(e['leading'])
                    pdf.setFont(e['font'],e['size']);pdf.setFillColor('#'+e['color'])
                    pdf.drawString(x*72,(H-y)*72-e['size']*.95-i*e['leading'],line)
            elif e['kind']=='image':
                ps.shapes.add_picture(e['path'],Inches(x),Inches(y),Inches(w),Inches(h))
                pdf.drawImage(e['path'],x*72,(H-y-h)*72,w*72,h*72)
            else:
                shape=ps.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e['radius'] else MSO_SHAPE.RECTANGLE,
                                          Inches(x),Inches(y),Inches(w),Inches(h))
                if e['radius']:shape.adjustments[0]=.03
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(e['fill'])
                if e['stroke']:
                    shape.line.color.rgb=RGBColor.from_string(e['stroke']);shape.line.width=Pt(.6)
                else:shape.line.fill.background()
                pdf.setFillColor('#'+e['fill']);pdf.setStrokeColor('#'+(e['stroke'] or e['fill']));pdf.setLineWidth(.6)
                if e['radius']:pdf.roundRect(x*72,(H-y-h)*72,w*72,h*72,3,fill=1,stroke=bool(e['stroke']))
                else:pdf.rect(x*72,(H-y-h)*72,w*72,h*72,fill=1,stroke=bool(e['stroke']))
        # Positive intersection of text boxes is unintended. Containers/images are exempt.
        for i,a in enumerate(texts):
            for b in texts[i+1:]:
                ix=min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x'])
                iy=min(a['y']+a['h'],b['y']+b['h'])-max(a['y'],b['y'])
                if ix>.015 and iy>.015:overlapping.append((s['index'],a['value'],b['value'],round(iy,3)))
        pdf.showPage()
    assert not overlapping,(t['id'],overlapping)
    pdf.save();ppt.save(folder/'guide.pptx')
    readback=Presentation(folder/'guide.pptx')
    assert len(readback.slides)==10 and all(s.notes_slide.notes_text_frame.text.strip() for s in readback.slides)
    with zipfile.ZipFile(folder/'guide.pptx') as z:assert z.testzip() is None
    reader=PdfReader(folder/'guide.pdf')
    assert len(reader.pages)==10
    for page,s in zip(reader.pages,slides):assert s['title'] in page.extract_text(),s['title']
    previews=folder/'preview';previews.mkdir(exist_ok=True)
    doc=pdfium.PdfDocument(str(folder/'guide.pdf'))
    for i in range(len(doc)):
        bitmap=doc[i].render(scale=1.5);im=bitmap.to_pil().convert('RGB')
        im.save(previews/f'{i+1:02d}.png');bitmap.close()
    copy=[s['copy'] for s in slides]
    stats=dict(slides=10,words=[len(re.findall(r'\S+',' '.join(v))) for v in copy],
               text_boxes=text_count,text_bounds='pass',text_collisions='none',
               pptx_zip='pass',notes='10/10',pdf_titles='10/10',pdf_rendered='10/10',
               copy_hash=hashlib.sha256(json.dumps(copy,ensure_ascii=False).encode()).hexdigest(),
               pdf_method='Same measured layout; not a native PowerPoint export')
    (folder/'validation.json').write_text(json.dumps(stats,indent=2)+'\n',encoding='utf-8')
    (folder/'speaker_notes.md').write_text('# DDRCV operator guide — '+t['name']+'\n\n'+
        '\n\n'.join(f'## {s["index"]:02d}. {s["title"]}\n\n{s["notes"]}' for s in slides)+'\n',encoding='utf-8')
    return stats


def comparisons(font_dir):
    font=ImageFont.truetype(str(font_dir/'segoeuib.ttf'),25)
    small=ImageFont.truetype(str(font_dir/'segoeui.ttf'),19)
    selected=[1,2,7]
    sheet=Image.new('RGB',(1800,1178),'#DDE3E9');draw=ImageDraw.Draw(sheet)
    for c,t in enumerate(THEMES):
        xx=c*600
        draw.text((xx+18,15),t['letter']+'  '+t['name'],font=font,fill='#182739')
        for r,n in enumerate(selected):
            im=Image.open(OUT/t['id']/'preview'/f'{n:02d}.png');im.thumbnail((590,332))
            yy=64+r*370;sheet.paste(im,(xx+5,yy))
            draw.text((xx+18,yy+337),f'Slide {n:02d}',font=small,fill='#35445A')
    sheet.save(OUT/'style_comparison.png')
    # Three high-resolution pages compare identical slides across the designs.
    pc=canvas.Canvas(str(OUT/'style_comparison.pdf'),pagesize=(1440,470))
    pc.setTitle('DDRCV — three visual directions')
    for n in selected:
        pc.setFillColor('#E8EDF2');pc.rect(0,0,1440,470,fill=1,stroke=0)
        pc.setFont('UIB',25);pc.setFillColor('#172B46')
        pc.drawString(20,434,f'Same content, three visual directions — slide {n:02d}')
        for r,t in enumerate(THEMES):
            xx=480*r+10
            pc.setFont('UIB',20);pc.drawString(xx+6,391,t['letter']+'  '+t['name'])
            pc.setFont('UI',15)
            words=t['description'].split();line='';lines=[]
            for word in words:
                test=(line+' '+word).strip()
                if pdfmetrics.stringWidth(test,'UI',15)>435:lines.append(line);line=word
                else:line=test
            lines.append(line)
            for j,line in enumerate(lines):pc.drawString(xx+6,91-j*21,line)
            pc.drawImage(str(OUT/t['id']/'preview'/f'{n:02d}.png'),xx,116,460,258.75)
        pc.showPage()
    pc.save()
    for t in THEMES:
        sheet=Image.new('RGB',(1200,5*366),'#DDE3E9');draw=ImageDraw.Draw(sheet)
        for i in range(10):
            im=Image.open(OUT/t['id']/'preview'/f'{i+1:02d}.png');im.thumbnail((592,333))
            x=(i%2)*600;y=(i//2)*366;sheet.paste(im,(x+4,y))
            draw.text((x+14,y+337),f'{i+1:02d}',font=small,fill='#24374B')
        sheet.save(OUT/t['id']/'contact_sheet.png')


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--font-dir',type=Path,default=Path('C:/Windows/Fonts'))
    args=parser.parse_args()
    for name,file in [('UI','segoeui.ttf'),('UIB','segoeuib.ttf'),('MonoB','consolab.ttf'),('Serif','georgia.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(args.font_dir/file)))
    OUT.mkdir(parents=True,exist_ok=True)
    stats={t['id']:render(Deck(t)) for t in THEMES}
    assert len({v['copy_hash'] for v in stats.values()})==1,'Variant copy differs'
    comparisons(args.font_dir)
    old=Presentation(ROOT/'doc/slides/ddrcv_operator_guide.pptx')
    old_counts=[len(re.findall(r'\S+',' '.join(sh.text for sh in s.shapes if sh.has_text_frame))) for s in old.slides]
    report=dict(variants=stats,identical_content=True,original_slides=len(old.slides),revised_slides=10,
                slide_reduction_percent=round((1-10/len(old.slides))*100,1),
                original_native_words_including_chrome=sum(old_counts),
                revised_content_words=sum(next(iter(stats.values()))['words']),
                revised_content_words_per_slide=next(iter(stats.values()))['words'])
    (OUT/'review_metrics.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,indent=2))


if __name__=='__main__':main()

"""Build editable PPTX, companion PDF, notes, and previews from one layout.

Requires python-pptx, reportlab, Pillow, pypdfium2, and pypdf. Uses Windows
Segoe UI fonts by default. See doc/slides/README.md for reproduction.
"""
from pathlib import Path
import argparse
import json
import math
import textwrap
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import pypdfium2 as pdfium
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'doc/slides'
ASSETS = OUT / 'assets'
W, H = 13.333333, 7.5
BG, PANEL, EDGE = '161C24', '222C38', '3C4C60'
WHITE, MUTED, BLUE, GREEN, AMBER = 'F3F6FA', 'B5C2D2', '83BFFF', '85DFBD', 'F4CA85'
SLIDES = []


def rect(s, x, y, w, h, fill=PANEL, stroke=None):
    s['elements'].append(dict(kind='rect', x=x, y=y, w=w, h=h, fill=fill, stroke=stroke))


def text(s, value, x, y, w, size=20, color=WHITE, bold=False, h=None):
    font = 'GuideBold' if bold else 'Guide'
    lines = []
    for paragraph in value.split('\n'):
        line = ''
        for word in paragraph.split():
            proposal = (line + ' ' + word).strip()
            if line and pdfmetrics.stringWidth(proposal, font, size) > w * 72:
                lines.append(line)
                line = word
            else:
                line = proposal
        lines.append(line)
    leading = size * 1.22
    height = len(lines) * leading / 72 + .035
    if h is not None and height > h:
        raise ValueError(f'Text exceeds allocation on slide {len(SLIDES)}: {value}')
    for line in lines:
        assert pdfmetrics.stringWidth(line, font, size) <= w * 72 + 1, line
    s['elements'].append(dict(kind='text', x=x, y=y, w=w, h=height,
                              lines=lines, size=size, color=color, bold=bold, leading=leading))
    return height


def slide(chapter, title, subtitle, notes):
    s = dict(chapter=chapter, title=title, notes=textwrap.dedent(notes).strip(), elements=[])
    SLIDES.append(s)
    text(s, chapter.upper(), .6, .35, 11.8, 11, BLUE, True)
    text(s, title, .6, .8, 12.1, 30, bold=True, h=.7)
    if subtitle:
        text(s, subtitle, .6, 1.42, 12.1, 15, MUTED, h=.56)
    return s


def card(s, x, y, w, h, title, body, accent=BLUE, size=19):
    rect(s, x, y, w, h)
    rect(s, x, y, .045, h, accent)
    th = text(s, title, x+.23, y+.2, w-.46, 22, accent, True)
    text(s, body, x+.23, y+.27+th, w-.46, size, h=h-.48-th)


def band(s, label, body, y=6.31, color=GREEN):
    rect(s, .6, y, 12.13, .58)
    text(s, label.upper(), .8, y+.13, 1.63, 12, color, True)
    text(s, body, 2.48, y+.115, 10.0, 15, h=.4)


def pic(s, name, x, y, w, h, caption=True):
    path = ASSETS / name
    with Image.open(path) as image:
        ratio = image.width / image.height
    pw, ph = w, w/ratio
    if ph > h:
        ph, pw = h, h*ratio
    s['elements'].append(dict(kind='image', path=str(path), x=x, y=y, w=pw, h=ph))
    if caption:
        text(s, 'Current app • isolated, stopped session • example paths',
             x, y+ph+.09, w, 10, MUTED)
    return pw, ph


def row(s, y, left, right, color=BLUE, x=.6, w=12.13, split=3.2, h=.74, size=18):
    rect(s, x, y, w, h)
    text(s, left, x+.2, y+.13, split-.3, size, color, True, h=h-.18)
    text(s, right, x+split+.15, y+.13, w-split-.35, size, h=h-.18)


def build_content():
    s = slide('Operator guide / Windows', 'DDRCV', '', '''
        Audience: the person setting up or running the tournament station.
        Walk through slides 1–16 before the event; use slides 17–21 at the desk.
        The screenshots are current desktop forms in an isolated stopped session,
        with example local paths and no credentials. They are not a live rehearsal.
        Source: doc/operator_walkthrough.md and doc/operator_reference.md.
        Source/UI snapshot: 18 September 2026. Recheck labels after application updates.
    ''')
    text(s, 'From first install\nto tournament operation', .6, 1.75, 8.2, 39, bold=True)
    text(s, 'Setup walkthrough + event-day reference', .63, 3.43, 10.8, 23, MUTED)
    for x, n, title, body in [( .6,'01','SET UP','Install and select your source'),
                              (4.73,'02','REHEARSE','Verify scores and integrations'),
                              (8.86,'03','OPERATE','Save, recover, and hand over')]:
        rect(s,x,4.45,3.87,1.57)
        text(s,n,x+.18,4.6,.65,26,GREEN,True)
        text(s,title,x+.95,4.67,2.7,15,BLUE,True)
        text(s,body,x+.2,5.23,3.47,18)
    text(s,'Desktop workflow • 18 September 2026 • Speaker notes included',.63,6.53,11.8,13,MUTED)

    s = slide('01 / Setup', 'What connects to what',
              'Start with recognition. Add the integrations your event uses.', '''
        DDRCV reads the full DDR World game feed and publishes live scores.
        ddr_score_render is a separate repository that displays those scores in a
        window. OBS can capture that window and optionally automate scene selection.
        Results capture saves evidence locally and can send it through a Discord webhook.
        Neither OBS nor the renderer is required to verify recognition itself.
        Source: doc/operator_walkthrough.md, introduction and sections 5–7.
    ''')
    card(s,.6,2.4,3.2,2.04,'Game feed','Capture device or RTSP\nFull 1280 × 720 frame')
    text(s,'→',3.99,3.01,.5,28,BLUE,True)
    card(s,4.66,2.4,3.3,2.04,'DDRCV','Recognize state\nRead and publish scores',GREEN)
    text(s,'→',8.17,3.01,.5,28,BLUE,True)
    card(s,8.92,2.05,3.81,1.8,'Score display','ddr_score_render → OBS',size=18)
    card(s,8.92,4.05,3.81,1.8,'Results evidence','Local screenshots\nOptional Discord delivery',size=18)
    band(s,'First milestone','A clean picture and accepted scores that match the machine.')

    s = slide('01 / Setup', 'Before you start',
              'Prepare the station and the event’s files.', '''
        Use the version or ZIP supplied by the event organizer. Extract it into a
        writable folder, not Program Files and not inside the ZIP. First setup needs
        internet. No existing Python install is necessary. Live capture is fixed at
        1280 × 720; a full-frame 16:9 recording can be resized explicitly.
        OBS, renderer assets, and a webhook are optional. Get the renderer repository
        and its assets from the organizer; DDRCV setup does not download them.
        Source: doc/operator_walkthrough.md sections 1–2; doc/installation.md.
    ''')
    card(s,.6,2.15,5.9,3.65,'Required','64-bit Windows + first-install internet\nDDRCV ZIP / checkout\nFull game feed at 1280 × 720\nOr a full-frame rehearsal recording',size=22)
    card(s,6.74,2.15,5.99,3.65,'If your event uses them','Saved event JSON profile\nOBS scene collection\nddr_score_render folder + assets\nDiscord channel webhook',GREEN,size=22)
    band(s,'Folder example',r'C:\Games\ddrcv  — extract the ZIP completely before running setup.')

    s = slide('01 / Setup', 'Install once, then use the launcher',
              'No preinstalled Python is required.', '''
        Double-click setup_desktop.bat. On a fresh install, Enter selects option 1:
        Install/use local Python in this checkout. Option 2 finds compatible installed
        Python versions; option 3 accepts a manual path. The supported interpreter is
        x64 CPython 3.12. On later runs, option 4 reuses the existing environment.
        Setup installs locked dependencies, provisions and tests actual OCR, and checks
        WebView2. Wait for Setup complete. If it fails, retain the exact error text.
        Do not move or delete .venv or .tools/python. Set up a fresh environment when
        moving to another computer. See doc/installation.md for troubleshooting.
    ''')
    for y,a,b in [(2.1,'1  Extract',r'Open the folder containing setup_desktop.bat.'),
                  (3.02,'2  Run setup','Double-click setup_desktop.bat; choose option 1.'),
                  (3.94,'3  Wait','Leave the console open until “Setup complete”.'),
                  (4.86,'4  Launch','Double-click launch_desktop.vbs.')]:
        row(s,y,a,b,h=.78)
    band(s,'Check','The desktop panel opens. Capture stays stopped until you click Start.')

    s = slide('01 / Setup', 'Find the four areas of the panel',
              'Start, Stop, Apply, and profile controls remain available across tabs.', '''
        This screenshot shows the real app with capture stopped. The UI may display
        unknown or absent players until a live session observes them.
        Status & configuration: source, score mode, health, names, accepted scores.
        OBS: connection, mapping, automation and manual scene switching.
        Renderer: repository path and launch controls.
        Results: screenshots, Discord, and recent results/recovery.
        Source: ddrcv/desktop/templates/index.html and static/desk.js.
    ''')
    pic(s,'panel.png',.6,2.06,7.14,4.75)
    for y,title,body in [(2.13,'Status & configuration','Choose the feed and inspect live reads.'),
                         (3.21,'OBS','Connect scenes and control automation.'),
                         (4.29,'Renderer','Launch the separate score display.'),
                         (5.37,'Results','Save screenshots, post, and recover.')]:
        text(s,title,8.04,y,4.65,21,BLUE,True)
        text(s,body,8.04,y+.4,4.65,18)

    s = slide('01 / Setup', 'Choose a clean source',
              'Status & configuration → Source', '''
        Capture camera: Scan, choose the device, then check Camera index. Device indices
        can differ between computers. RTSP stream: enter the URL; frames must be 720p.
        Local video: Browse to the file, enabling Resize full 16:9 file to 720p if needed.
        Borders, cropped broadcasts, overlays, or stretching move the recognition regions.
        Still image diagnostics cannot establish temporal score or state confirmation.
        Leave publisher localhost:9000 for a single-computer setup. Run only one app,
        workbench, or CLI against that publisher port and Results outbox.
        Source: doc/operator_walkthrough.md section 3.
    ''')
    pic(s,'source.png',.6,2.05,12.13,3.03)
    text(s,'LIVE CAMERA',.6,5.48,3.7,13,BLUE,True)
    text(s,'Scan → select device\nCheck Camera index',.6,5.79,3.7,19)
    text(s,'RTSP',4.75,5.48,3.65,13,BLUE,True)
    text(s,'Enter Stream URL\nRequire 1280 × 720',4.75,5.79,3.65,19)
    text(s,'RECORDING',8.86,5.48,3.86,13,BLUE,True)
    text(s,'Browse to Local video\nResize only full 16:9 files',8.86,5.79,3.86,19)

    s = slide('01 / Setup', 'Match the score mode to the game',
              'This setting affects gameplay filtering and results interpretation.', '''
        Select EX when the game shows the event EX score. Select Money for the ordinary
        score toward 1,000,000. Desktop defaults to EX; the browser workbench defaults
        to Money. A profile can override either, so always verify the selected mode.
        A wrong score mode may reject otherwise valid score reads and uses the wrong
        results crop. Leave Custom filter to the profile maintainer.
        Money results do not declare an EX winner.
        Source: driver_config.py; doc/results_publishing.md, confirmation/mode rules.
    ''')
    card(s,.6,2.16,5.9,3.61,'EX score','For the event’s EX scoring.\nDesktop default.\nVerify with an actual song.',BLUE,size=24)
    card(s,6.74,2.16,5.99,3.61,'Money score','For the score toward 1,000,000.\nBrowser workbench default.\nVerify after switching apps.',GREEN,size=24)
    band(s,'Remember','An old profile may select a different mode. Check before pressing Start.',color=AMBER)

    s = slide('02 / Rehearsal', 'Prove recognition before adding integrations',
              'Click Start, then observe a complete song from its splash.', '''
        The screenshot is a stopped UI for orientation, not proof of a running capture.
        During your test, look for Running, healthy capture, a confirmed state following
        the game, and recent accepted reads that agree with the machine.
        View stream fetches one snapshot. Refresh frame fetches the next; it is not a
        continuous preview. Starting mid-song can lack player-presence evidence.
        The final score, including a full-combo ending, must be compared with the machine.
        A briefly delayed confirmed state is expected because recognition needs multiple
        observations. Source: operator_walkthrough.md section 4; tournament_desktop.md.
    ''')
    pic(s,'status.png',.6,2.05,12.13,3.0)
    for x,title,body in [( .6,'1  Picture','View stream → Refresh frame.\nCheck full, clean framing.'),
                         (4.75,'2  State','Song splash → gameplay → results.\nConfirmed state follows the game.'),
                         (8.86,'3  Scores','Compare accepted reads and ages.\nVerify the final score too.')]:
        text(s,title,x,5.35,3.85,21,GREEN,True)
        text(s,body,x,5.84,3.85,18)

    s = slide('02 / Rehearsal · optional', 'Add the score renderer',
              'Renderer → select repository → apply → start or launch', '''
        Obtain the separate ddr_score_render repository and assets from the organizer.
        Select its folder with Browse; the folder must contain apps/tbd5.py.
        Set Launch on Start if desired. Apply configuration while stopped, then Start.
        If already running, Apply and Restart resets tracking, so use it between songs.
        With automatic launch off, use Launch renderer after starting capture.
        Check the actual window, then capture it in OBS. Publisher clients or a running
        process alone do not establish visible output. Buttons use the applied path/port.
        Source: doc/operator_walkthrough.md section 5; tournament_integrations.py.
    ''')
    pic(s,'renderer.png',.6,2.05,12.13,2.36)
    card(s,.6,4.78,3.88,1.39,'Choose the folder','Separate ddr_score_render repo.',size=16)
    card(s,4.73,4.78,3.88,1.39,'Apply the path','Then Start / Launch renderer.',size=16)
    card(s,8.86,4.78,3.87,1.39,'Inspect the window','Compare scores; capture in OBS.',GREEN,size=16)
    band(s,'Check','Visible scores agree with DDRCV’s accepted scores.')

    s = slide('02 / Rehearsal · optional', 'Connect OBS before enabling automation',
              'Use exact scene names that already exist in the event’s OBS collection.', '''
        Enable OBS's WebSocket server. Enter the host, port and password into DDRCV;
        usual same-computer settings are localhost and 4455. Check Connect on Start.
        Keep Enable scene automation on Start off for the initial rehearsal.
        Edit the JSON map, retaining quotes and commas. Apply configuration and Start;
        Connect if necessary, then Check scenes. Enable automation after validation.
        song_select, login and total_result default to Commentator; song_splash,
        song_playing and song_result default to Gameplay. Null/omitted mappings hold;
        unknown always holds. No scenes are created by DDRCV.
        Source: doc/operator_walkthrough.md section 6; publish/obs_policy.py.
    ''')
    pic(s,'obs.png',.6,2.02,7.18,4.76)
    text(s,'1  Enter connection details',8.03,2.17,4.68,22,BLUE,True)
    text(s,'Host, port, password.\nApply settings before Connect.',8.03,2.94,4.6,19)
    text(s,'2  Check the mapping',8.03,3.89,4.68,22,BLUE,True)
    text(s,'Use exact scene names.\nClick Check scenes.',8.03,4.37,4.6,19)
    text(s,'3  Enable automation',8.03,5.29,4.68,22,GREEN,True)
    text(s,'Watch the actual OBS scene\nthrough a complete song.',8.03,5.79,4.6,19)

    s = slide('02 / Rehearsal · optional', 'Practice manual OBS control',
              'Know how to pause automation before the event needs it.', '''
        In DDRCV's OBS tab, select an existing scene and use Switch scene & pause
        automation. Resume with Enable automation; it reconciles to the latest
        confirmed state. If changing a scene directly in OBS, pause DDRCV automation
        first so it does not change the scene back. This does not stop recognition.
        After a successful rehearsal, choose whether automation should enable on Start,
        then apply and save the profile. Source: tournament_desktop.md and obs_policy.py.
    ''')
    card(s,.6,2.28,5.9,3.29,'Take manual control','Choose a scene.\nClick “Switch scene & pause automation”.\nVerify the scene in OBS.',BLUE,size=23)
    card(s,6.74,2.28,5.99,3.29,'Resume automatic control','Click “Enable automation”.\nVerify the scene matches the confirmed game state.',GREEN,size=23)
    band(s,'Direct OBS edit','Pause DDRCV automation before changing a scene inside OBS.',color=AMBER)

    s = slide('02 / Rehearsal', 'Save results screenshots',
              'Results → Capture results → Screenshot folder → Apply results settings', '''
        New profiles enable Capture results, but an existing saved off value is honored.
        Choose a writable screenshot folder. Default copies are under screenshots in
        the checkout. Durable originals and delivery records live in .workbench/results
        unless configured otherwise. Changing the screenshot folder does not move old files.
        Filenames use UTC time and a unique event suffix. Parsed names can be included
        only when recognized. Apply results settings does not restart capture, clear
        gameplay scores, or reconnect OBS/renderer. Save the JSON profile separately.
        Source: operator_walkthrough.md section 7; results_publishing.md.
    ''')
    pic(s,'results-capture.png',.6,2.01,8.45,4.72)
    text(s,'Capture results',9.33,2.22,3.38,22,BLUE,True)
    text(s,'Check it even when using an existing profile.',9.33,2.73,3.38,20)
    text(s,'Screenshot folder',9.33,3.86,3.38,22,BLUE,True)
    text(s,'Choose a writable place you can find later.',9.33,4.36,3.38,20)
    text(s,'Apply results settings',9.33,5.49,3.38,21,GREEN,True)
    text(s,'Keeps gameplay scores.',9.33,6.23,3.38,18)

    s = slide('02 / Rehearsal · optional', 'Test Discord, then enable automatic posts',
              'A webhook test sends a real message to the selected channel.', '''
        Ask the organizer/channel administrator for the intended results webhook. Enter
        it in Results. Test Discord Webhook uses the URL currently entered even when
        capture and automatic publishing are off. It does not apply pending settings
        or enable publishing. Verify the channel before enabling Publish results to Discord.
        Enabling publishing also checks Capture results. Click Apply results settings.
        Complete a test song, inspect Recent results & recovery, then inspect Discord.
        During a local rehearsal leave publishing off, or use the organizer's test channel.
        Source: desktop/templates/index.html; operator_walkthrough.md section 7.
    ''')
    pic(s,'discord.png',.6,2.06,12.13,2.46)
    for x,n,title,body in [( .6,'1','Enter + test','Webhook URL → Test Discord Webhook.'),
                           (4.75,'2','Verify channel','Find the actual test message.'),
                           (8.86,'3','Enable + apply','Publish results to Discord → Apply results settings.')]:
        text(s,n+'  '+title,x,4.94,3.83,22,GREEN,True)
        text(s,body,x,5.54,3.83,20)

    s = slide('02 / Rehearsal', 'Choose how much results text to read',
              'These choices preserve live gameplay scoring.', '''
        Screenshots only and Read names on results apply immediately, without a general
        restart. Screenshots only skips both results score and name OCR. The detected
        results episode saves on departure or after its original ten-second deadline.
        Capture results must stay enabled; Discord must also be enabled for posts.
        With parsing enabled, confirmed scores can post before names; later confirmed
        names update the same message. With name parsing disabled, scores use P1/P2.
        The login/song splash name switch is separate and applies immediately.
        These switches do not correct saved results or recover an entirely missed screen.
        Source: results_publishing.md and tournament_desktop.md.
    ''')
    row(s,2.14,'Scores + names','Screenshots only OFF; Read names on results ON.',h=.96)
    row(s,3.3,'Scores + P1/P2','Screenshots only OFF; Read names on results OFF.',h=.96)
    row(s,4.46,'Screenshot evidence','Screenshots only ON. Save on departure or after 10 seconds.',GREEN,h=.96)
    band(s,'Separate control','Live-session names use “Read names on login/song splash”.')

    s = slide('02 / Rehearsal', 'Save a profile and understand Apply',
              'The form, running session, and saved file can contain different settings.', '''
        Open loads a profile into the form without starting capture. Save/Save as writes
        the form into a JSON profile without applying pending changes to the current run.
        Start uses the current form. Apply configuration works while stopped; Start runs it.
        Apply and Restart during a run confirms a reset of live score/state tracking.
        Use Results-only Apply for those controls; name and screenshot-only switches are
        immediate. Save a named profile after changing settings, even when the local draft
        is automatically retained. Profiles contain real secrets and absolute paths.
        Source: desktop/controller.py, static/desk.js; operator_walkthrough.md section 8.
    ''')
    card(s,.6,2.14,3.88,3.73,'Save / Save as…','Write the current form to a named JSON profile.\n\nRunning settings stay as they were.',BLUE,size=22)
    card(s,4.73,2.14,3.88,3.73,'Start','Start recognition using the current form settings.\n\nUse after opening the event profile.',GREEN,size=22)
    card(s,8.86,2.14,3.87,3.73,'Apply and Restart','Apply pending changes during a run.\n\nResets live score tracking.',AMBER,size=22)
    band(s,'Between songs','Use the full Apply and Restart for source/mode/integration changes.',color=AMBER)

    s = slide('02 / Rehearsal', 'Finish with a real event rehearsal',
              'Installation checks do not establish tournament readiness.', '''
        Use the actual tournament computer, capture device, renderer and OBS scene
        collection. Observe select → splash → gameplay → song results → total results.
        Validate full-combo final scores, single-player presence and starting mid-song.
        Exercise capture reconnect, manual scene selection/resume, and window close/reopen.
        Check actual Results files and Discord messages, including recovery from delivery
        failure. If a results screen is never detected, automatic recovery has no episode.
        Preserve the event's recording or manual evidence procedure.
        Source: operator_walkthrough.md section 8; tournament_handoff.md section 5.
    ''')
    for y,left,right in [(2.13,'Game flow','Selection → splash → gameplay → results.'),
                         (3.02,'Score accuracy','Two players, one player, and a full-combo ending.'),
                         (3.91,'Connections','Capture reconnect; manual OBS scene and resume.'),
                         (4.80,'Evidence','Screenshot files, Discord post, and recovery actions.')]:
        row(s,y,left,right,h=.77)
    band(s,'Ready when','Another operator can repeat startup and recovery using the saved profile.')

    s = slide('03 / Event-day reference', 'Start the station in five steps',
              'Keep this slide available beside the event profile.', '''
        If the panel was closed but the backend remains active, reopening it does not
        create another recognition session. Fully exiting or a backend crash requires
        relaunch and an explicit Start. Profile restoration never starts capture itself.
        Verify mode and capture index every time hardware or computers change. A saved
        profile can restore automatic OBS/Discord actions; inspect it before Start.
        For problems, check fresh capture and accepted-read ages independently.
        Source: operator_reference.md, Start an event session.
    ''')
    for y,left,right in [(2.07,'1  Connect','Capture device ready; OBS open if used.'),
                         (2.9,'2  Launch','launch_desktop.vbs → Open… event profile if needed.'),
                         (3.73,'3  Verify','Source, EX/Money, Results, and integration settings.'),
                         (4.56,'4  Start','Running + healthy capture; inspect View stream.'),
                         (5.39,'5  Confirm','Accepted scores, actual score display, scenes, results post.')]:
        row(s,y,left,right,h=.72,size=18)

    s = slide('03 / Event-day reference', 'Recover a saved result',
              'Results → Recent results & recovery', '''
        Inspect the saved screenshot before sending. Send/retry requires capture Running,
        Capture results enabled, Discord publishing enabled, and a configured webhook.
        Saved: use Send saved result. Failed: fix the issue and Retry delivery.
        Queued/Sending/Retrying: wait and watch connectivity. Sent: Open Discord message.
        Retry name update edits the existing message after a failed name update.
        Republish under More recovery options confirms creation of another Discord post.
        Disabling publishing and applying pauses queued deliveries; in-flight sends may
        finish. Re-enabling resumes queued work. A completely missed screen cannot recover.
        Source: operator_reference.md, Results recovery; static/desk.js.
    ''')
    for y,left,right in [(2.07,'Saved','Send saved result'),
                         (2.84,'Queued / Sending / Retrying','Watch the status and connection.'),
                         (3.61,'Failed','Fix the cause → Retry delivery.'),
                         (4.38,'Sent','Open Discord message; Retry name update if offered.'),
                         (5.15,'Need another post?','More recovery options → Republish creates a duplicate.')]:
        row(s,y,left,right,split=4.0,h=.67,size=16)
    band(s,'Required','Capture running + Capture results + Discord enabled + webhook.',color=AMBER)

    s = slide('03 / Event-day reference', 'When something looks wrong',
              'Check the failing part before restarting the whole session.', '''
        Preview is a requested still; Refresh frame updates it. Capture freshness and
        accepted-read age differ: frames can be fresh while score reads are rejected.
        Check complete 1280x720 framing, source, score mode, and read rejection reasons.
        If a wrongly high accepted score is stuck, use Reset P1/P2 tracking after checking
        the input. It clears only that live score and does not edit saved results.
        A camera switch tests the new source before replacing the old; success resets
        tracking. If a driver cannot open both, Stop/Start the new source between songs.
        Port/outbox conflicts: stop the other desktop/workbench/CLI; a closed panel can
        still own the session. See operator_reference.md for the full troubleshooting table.
    ''')
    for y,left,right in [(2.07,'Frozen preview','Refresh frame; check capture age and recent errors.'),
                         (2.88,'Wrong or stuck score','Check framing + mode; inspect rejection details.'),
                         (3.69,'One bad accepted score','Read details & recent errors → Reset P1/P2 tracking.'),
                         (4.5,'Blank renderer','Check publisher clients, path, status, and actual window.'),
                         (5.31,'Wrong OBS scene','Check connection, scene names, and automation state.')]:
        row(s,y,left,right,split=3.85,h=.7,size=17)
    text(s,'Results text unreliable? Enable Screenshots only; gameplay scoring continues.',.8,6.46,11.9,17,GREEN)

    s = slide('03 / Event-day reference', 'Closing the panel keeps capture running',
              'Use Exit application… when the station should actually stop.', '''
        Window close/Alt+F4 closes the disposable UI, while the backend keeps capture,
        publishing, automation and owned renderer active. Open panel from the tray or
        relaunch the VBS to reopen without resetting scores. Exit application confirms
        stopping the backend and its owned renderer; OBS itself remains open.
        At session end inspect unresolved deliveries, save intended profile changes,
        then exit. Keep the screenshot copies and .workbench/results outbox for evidence.
        Desktop settings/logs: %LOCALAPPDATA%/DDRCV/tournament. For updates fully exit,
        update source, rerun setup using option 4 for a compatible environment, and rehearse.
        Source: operator_reference.md, Files and shutdown; tournament_desktop.md.
    ''')
    card(s,.6,2.2,5.9,3.27,'Close window / Alt+F4','Capture keeps running.\nReopen from the tray or launch_desktop.vbs.\nLive scores remain intact.',BLUE,size=23)
    card(s,6.74,2.2,5.99,3.27,'Exit application…','Stops the application after confirmation.\nOBS itself stays open.\nNext launch needs Start.',AMBER,size=23)
    band(s,'Before exiting','Check unresolved deliveries → save the profile → retain event evidence.')

    s = slide('03 / Handoff', 'Leave the next operator a complete handoff',
              'Fill these in for the actual tournament computer.', '''
        Fill the event setup table in doc/operator_walkthrough.md before handover.
        Include DDRCV folder/version, saved JSON path, source device/index/resolution,
        EX/Money, renderer folder, OBS scene collection/mapping, screenshot folder,
        Discord channel/contact and manual evidence procedure.
        Do not write webhook tokens or passwords into this presentation; saved JSON
        profiles contain actual credentials despite masked UI fields. Share them only
        with authorized operators. Verify absolute paths and camera index on each computer.
        Further reading: doc/operator_walkthrough.md, doc/operator_reference.md,
        doc/installation.md, doc/tournament_handoff.md, doc/results_publishing.md.
    ''')
    card(s,.6,2.12,5.9,2.96,'Station details','Profile path + DDRCV folder/version\nCapture device / index + score mode\nRenderer folder + OBS scenes',size=21)
    card(s,6.74,2.12,5.99,2.96,'Evidence and support','Screenshot folder + Discord channel\nTechnical contact\nRecording / manual-results procedure',GREEN,size=21)
    text(s,'Full walkthrough',.63,5.44,3.6,17,BLUE,True)
    text(s,'doc/operator_walkthrough.md',4.15,5.44,8.45,18)
    text(s,'Quick reference',.63,5.92,3.6,17,BLUE,True)
    text(s,'doc/operator_reference.md',4.15,5.92,8.45,18)
    text(s,'Keep credential-bearing profiles private; verify paths on each computer.',.63,6.52,12.05,16,AMBER)


def render(font_dir):
    OUT.mkdir(parents=True, exist_ok=True)
    pdfmetrics.registerFont(TTFont('Guide', str(font_dir / 'segoeui.ttf')))
    pdfmetrics.registerFont(TTFont('GuideBold', str(font_dir / 'segoeuib.ttf')))
    build_content()
    ppt = Presentation()
    ppt.slide_width, ppt.slide_height = Inches(W), Inches(H)
    ppt.core_properties.title = 'DDRCV — Setup and Operator Guide'
    ppt.core_properties.subject = 'Windows installation, rehearsal, and event operation'
    ppt.core_properties.author = 'DDRCV'
    pdf = canvas.Canvas(str(OUT / 'ddrcv_operator_guide.pdf'), pagesize=(W*72,H*72))
    pdf.setTitle(ppt.core_properties.title)
    pdf.setAuthor('DDRCV')
    notes = ['# DDRCV operator guide — speaker notes\n',
             'Slide content and screenshots reflect the repository on 18 September 2026.\n']
    element_count = 0
    for idx, s in enumerate(SLIDES, 1):
        rect(s,.6,7.06,12.13,.012,EDGE)
        rect(s,.6,7.06,12.13*idx/len(SLIDES),.022,BLUE)
        text(s,'DDRCV  /  '+s['chapter'],.6,7.17,10.9,9,MUTED)
        text(s,f'{idx:02d} / {len(SLIDES):02d}',11.78,7.16,.95,10,MUTED)
        ps = ppt.slides.add_slide(ppt.slide_layouts[6])
        ps.background.fill.solid()
        ps.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        ps.notes_slide.notes_text_frame.text = s['notes']
        pdf.bookmarkPage(f'slide-{idx}')
        pdf.addOutlineEntry(f'{idx:02d}. {s["title"]}', f'slide-{idx}', level=0)
        pdf.setFillColor('#'+BG)
        pdf.rect(0,0,W*72,H*72,stroke=0,fill=1)
        for e in s['elements']:
            assert e['x'] >= 0 and e['y'] >= 0
            assert e['x']+e['w'] <= W+.01, (idx,e)
            assert e['y']+e['h'] <= H+.01, (idx,e)
            element_count += 1
            x,y,w,h = [e[v] for v in ('x','y','w','h')]
            if e['kind']=='rect':
                shape = ps.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
                shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(e['fill'])
                shape.line.fill.background()
                pdf.setFillColor('#'+e['fill']); pdf.rect(x*72,(H-y-h)*72,w*72,h*72,fill=1,stroke=0)
            elif e['kind']=='image':
                ps.shapes.add_picture(e['path'],Inches(x),Inches(y),width=Inches(w),height=Inches(h))
                pdf.drawImage(e['path'],x*72,(H-y-h)*72,w*72,h*72)
            else:
                shape = ps.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                tf = shape.text_frame
                tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
                tf.word_wrap=False; tf.auto_size=MSO_AUTO_SIZE.NONE
                for i,line in enumerate(e['lines']):
                    p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
                    p.text=line
                    p.font.name='Segoe UI'; p.font.size=Pt(e['size']); p.font.bold=e['bold']
                    p.font.color.rgb=RGBColor.from_string(e['color'])
                    p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(e['leading'])
                    pdf.setFillColor('#'+e['color'])
                    pdf.setFont('GuideBold' if e['bold'] else 'Guide',e['size'])
                    pdf.drawString(x*72,(H-y)*72-e['size']*.95-i*e['leading'],line)
        pdf.showPage()
        notes.append(f"\n## {idx:02d}. {s['title']}\n\n{s['notes']}\n")
    pdf.save()
    ppt_path = OUT / 'ddrcv_operator_guide.pptx'
    ppt.save(ppt_path)
    (OUT/'speaker_notes.md').write_text('\n'.join(notes),encoding='utf-8')
    # Independent package and PDF checks; no runtime/application tests required.
    readback = Presentation(ppt_path)
    assert len(readback.slides)==len(SLIDES)
    assert all(x.notes_slide.notes_text_frame.text.strip() for x in readback.slides)
    with zipfile.ZipFile(ppt_path) as package:
        assert package.testzip() is None
    reader = PdfReader(OUT/'ddrcv_operator_guide.pdf')
    assert len(reader.pages)==len(SLIDES)
    for s,p in zip(SLIDES, reader.pages):
        assert s['title'] in p.extract_text(),s['title']
    document = pdfium.PdfDocument(str(OUT/'ddrcv_operator_guide.pdf'))
    previews = OUT/'preview'
    previews.mkdir(exist_ok=True)
    thumbs=[]
    for i in range(len(document)):
        bitmap = document[i].render(scale=1.5)
        im=bitmap.to_pil().convert('RGB')
        im.save(previews/f'slide-{i+1:02d}.png')
        im.thumbnail((480,270))
        thumbs.append(im.copy())
        bitmap.close()
    sheet=Image.new('RGB',(480*3,302*math.ceil(len(thumbs)/3)), '#E4E9F0')
    draw=ImageDraw.Draw(sheet)
    font=ImageFont.truetype(str(font_dir/'segoeui.ttf'),15)
    for i,im in enumerate(thumbs):
        x,y=(i%3)*480,(i//3)*302
        sheet.paste(im,(x,y));draw.text((x+10,y+275),f'{i+1:02d}  {SLIDES[i]["title"]}',font=font,fill='#162030')
    sheet.save(OUT/'contact_sheet.png')
    report=dict(slides=len(SLIDES),notes_slides=len(readback.slides),elements_checked=element_count,
                bounds='pass',text_width_and_allocated_height='pass',pptx_zip='pass',
                pdf_page_titles='pass',pdf_pages_rendered=len(document),
                ui_images='Real desktop forms; isolated stopped session; no credentials or integrations',
                pdf_method='Companion PDF drawn from the same layout specification; not a PowerPoint export')
    (OUT/'validation.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--font-dir',type=Path,default=Path('C:/Windows/Fonts'))
    args=parser.parse_args()
    render(args.font_dir)
